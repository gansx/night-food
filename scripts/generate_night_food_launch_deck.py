from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("docs/night-food-launch-deck.pptx")

WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

BG = RGBColor(2, 5, 7)
PANEL = RGBColor(7, 17, 26)
PANEL_2 = RGBColor(11, 27, 34)
TEXT = RGBColor(245, 232, 220)
MUTED = RGBColor(169, 148, 131)
BRAND = RGBColor(239, 159, 93)
WARN = RGBColor(232, 119, 72)
BLUE = RGBColor(34, 116, 141)
GREEN = RGBColor(41, 90, 68)
LINE = RGBColor(231, 201, 174)


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def blank(prs: Presentation):
    return prs.slide_layouts[6]


def fill_bg(slide, accent: RGBColor = BRAND):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDE_W, WIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    # Stage glow
    for idx, (x, y, size, color) in enumerate(
        [
            (Inches(-1.2), Inches(-0.8), Inches(4.9), WARN),
            (Inches(9.6), Inches(-1.1), Inches(5.4), BLUE),
            (Inches(5.0), Inches(5.2), Inches(5.8), GREEN),
        ]
    ):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
        dot.fill.solid()
        dot.fill.fore_color.rgb = color if idx != 0 else accent
        dot.fill.transparency = 72
        dot.line.fill.background()

    # Fine grid lines
    for i in range(0, 14):
        x = Inches(i)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, 0, x, WIDE_H)
        line.line.color.rgb = rgb("1d2a2f")
        line.line.transparency = 55
        line.line.width = Pt(0.35)
    for i in range(0, 8):
        y = Inches(i)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, y, WIDE_W, y)
        line.line.color.rgb = rgb("1d2a2f")
        line.line.transparency = 55
        line.line.width = Pt(0.35)


def set_text(run, size=24, bold=False, color=TEXT, font="Microsoft YaHei"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def textbox(slide, text, x, y, w, h, size=24, bold=False, color=TEXT, align=PP_ALIGN.LEFT, font="Microsoft YaHei"):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_text(run, size=size, bold=bold, color=color, font=font)
    return box


def title(slide, text, subtitle=None):
    textbox(slide, text, Inches(0.78), Inches(0.72), Inches(8.5), Inches(0.72), size=28, bold=True, color=TEXT)
    if subtitle:
        textbox(slide, subtitle, Inches(0.8), Inches(1.34), Inches(8.0), Inches(0.35), size=11, color=MUTED)


def kicker(slide, text, x=0.78, y=0.52, w=2.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.36))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb("251b10")
    shape.line.color.rgb = BRAND
    shape.line.transparency = 35
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text.upper()
    set_text(r, size=9, bold=True, color=TEXT, font="Aptos")
    return shape


def card(slide, x, y, w, h, heading, body=None, color=PANEL, accent=BRAND, number=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = 5
    shape.line.color.rgb = accent
    shape.line.transparency = 58
    shape.line.width = Pt(1.0)

    if number:
        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18), Inches(y + 0.18), Inches(0.38), Inches(0.38))
        n.fill.solid()
        n.fill.fore_color.rgb = accent
        n.line.fill.background()
        nf = n.text_frame
        nf.clear()
        nf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = nf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(number)
        set_text(r, size=9, bold=True, color=BG, font="Aptos")
        tx = x + 0.68
        tw = w - 0.9
    else:
        tx = x + 0.24
        tw = w - 0.48

    textbox(slide, heading, Inches(tx), Inches(y + 0.22), Inches(tw), Inches(0.36), size=15, bold=True)
    if body:
        textbox(slide, body, Inches(x + 0.25), Inches(y + 0.74), Inches(w - 0.5), Inches(h - 0.8), size=10.5, color=MUTED)
    return shape


def bullet_list(slide, items, x, y, w, h, size=18, color=TEXT, gap=0.42):
    for idx, item in enumerate(items):
        yy = y + idx * gap
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.08), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = BRAND if idx % 2 == 0 else BLUE
        dot.line.fill.background()
        textbox(slide, item, Inches(x + 0.25), Inches(yy), Inches(w - 0.25), Inches(0.3), size=size, color=color)


def footer(slide, idx):
    textbox(slide, "Night Food / 家宴星球", Inches(0.78), Inches(7.04), Inches(3), Inches(0.25), size=8, color=MUTED, font="Aptos")
    textbox(slide, f"{idx:02d}", Inches(12.05), Inches(7.0), Inches(0.5), Inches(0.28), size=10, color=MUTED, align=PP_ALIGN.RIGHT, font="Aptos")


def connector(slide, x1, y1, x2, y2, color=BRAND):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.25)
    line.line.transparency = 15
    return line


def add_slide(prs, idx, accent=BRAND):
    slide = prs.slides.add_slide(blank(prs))
    fill_bg(slide, accent)
    footer(slide, idx)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H

    # 01 Cover
    s = add_slide(prs, 1, WARN)
    kicker(s, "Family Web3 Home", 0.82, 0.72, 2.35)
    textbox(s, "家宴星球", Inches(0.75), Inches(1.35), Inches(6.8), Inches(1.2), size=60, bold=True)
    textbox(s, "Night Food", Inches(0.82), Inches(2.48), Inches(3.4), Inches(0.42), size=20, color=BRAND, bold=True, font="Aptos")
    textbox(s, "把点餐、任务和积分，变成一个家庭协作的小宇宙。", Inches(0.82), Inches(3.22), Inches(6.9), Inches(0.48), size=20, color=MUTED)
    card(s, 8.55, 1.1, 3.55, 4.75, "产品完成版", "Web 家庭端 / Admin 家主管理端\nCloudflare 云部署\nSupabase 数据与认证\nPlaywright 线上回归", color=rgb("07111a"), accent=BRAND)
    for i, size in enumerate([4.2, 3.15, 2.1]):
        orbit = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.8 + i * 0.35), Inches(0.7 + i * 0.5), Inches(size), Inches(size))
        orbit.fill.background()
        orbit.line.color.rgb = LINE
        orbit.line.transparency = 78
        orbit.line.width = Pt(1)

    # 02 Problem
    s = add_slide(prs, 2, BLUE)
    title(s, "家庭每天都在协作，只是缺少一个好界面", "从聊天记录、口头提醒到可持续的家庭协作闭环")
    bullet_list(
        s,
        ["今天吃什么，经常散落在聊天记录里。", "家务任务靠口头提醒，反馈感很弱。", "奖励和积分没有体系，参与感难以持续。", "家主需要维护规则，但不应该背负管理焦虑。"],
        0.95,
        2.0,
        6.1,
        2.3,
        size=18,
        gap=0.62,
    )
    card(s, 8.0, 1.85, 3.9, 2.9, "设计目标", "让家庭日常变得清楚、轻量、有反馈，而不是把家庭变成公司。", color=PANEL_2, accent=BLUE)

    # 03 Position
    s = add_slide(prs, 3, BRAND)
    title(s, "产品定位", "家庭协作式点餐与积分任务平台")
    textbox(s, "Night Food 不是餐饮 SaaS，\n而是家庭内部的小型协作网络。", Inches(0.9), Inches(1.95), Inches(6.6), Inches(1.25), size=30, bold=True)
    card(s, 0.95, 4.35, 2.55, 1.25, "家人点餐", "点餐和订单状态", accent=BRAND, number=1)
    card(s, 3.75, 4.35, 2.55, 1.25, "家主维护", "菜单、任务、规则", accent=BLUE, number=2)
    card(s, 6.55, 4.35, 2.55, 1.25, "任务积分", "完成任务赚积分", accent=WARN, number=3)
    card(s, 9.35, 4.35, 2.55, 1.25, "家庭闭环", "积分回到点餐", accent=GREEN, number=4)

    # 04 Two apps
    s = add_slide(prs, 4, GREEN)
    title(s, "两个入口，两种角色", "家人要简单，家主要掌控")
    card(s, 0.9, 1.8, 5.25, 3.65, "Web 家庭端", "点餐\n订单\n任务\n我的\n家庭加入", color=PANEL, accent=BRAND)
    card(s, 7.15, 1.8, 5.25, 3.65, "Admin 家主管理端", "菜单\n订单\n任务\n积分\n成员\n规则", color=PANEL, accent=BLUE)
    connector(s, 6.15, 3.55, 7.15, 3.55, BRAND)
    textbox(s, "同一套家庭数据，两个清晰入口", Inches(3.9), Inches(6.0), Inches(5.8), Inches(0.35), size=18, color=MUTED, align=PP_ALIGN.CENTER)

    # 05 Loop
    s = add_slide(prs, 5, BRAND)
    title(s, "核心闭环", "菜单 -> 点餐 -> 订单 -> 任务 -> 积分 -> 再点餐")
    nodes = [
        (1.0, 2.2, "菜单", "家主维护"),
        (3.05, 1.35, "点餐", "家人选择"),
        (5.35, 2.2, "订单", "处理状态"),
        (7.65, 1.35, "任务", "领取提交"),
        (9.95, 2.2, "审核", "家主确认"),
        (5.35, 4.65, "积分", "消费与奖励"),
    ]
    for i, (x, y, h, b) in enumerate(nodes, 1):
        card(s, x, y, 1.65, 1.05, h, b, accent=[BRAND, BLUE, WARN, GREEN, BRAND, BLUE][i - 1], number=i)
    for a, b in [(1.85, 3.05), (3.9, 5.35), (6.2, 7.65), (8.5, 9.95)]:
        connector(s, a, 2.75 if a != 3.9 else 1.85, b, 2.75 if b != 7.65 else 1.85)
    connector(s, 10.75, 3.25, 6.2, 4.95, WARN)
    connector(s, 5.35, 5.2, 1.82, 3.25, BLUE)

    # 06 Family web
    s = add_slide(prs, 6, WARN)
    title(s, "家庭端体验", "给家人用的，不是给管理员用的")
    bullet_list(s, ["看图点餐，只显示菜名、图片和积分。", "订单状态清楚，允许范围内可取消。", "领取任务，提交完成，获得积分反馈。", "移动端底部导航，日常使用更顺手。"], 0.95, 1.85, 6.7, 2.8, size=18, gap=0.62)
    card(s, 8.15, 1.65, 3.65, 3.95, "Mobile First", "家庭场景大多数发生在手机上。\n界面以触控、卡片和底部导航为核心。", color=PANEL_2, accent=WARN)

    # 07 Admin
    s = add_slide(prs, 7, BLUE)
    title(s, "家主管理台", "家主不是老板，是家庭规则的维护者")
    for i, (x, y, h, b, c) in enumerate(
        [
            (0.9, 1.65, "菜单管理", "分类、菜品、图片、推荐", BRAND),
            (4.55, 1.65, "订单管理", "确认、制作、完成、取消", BLUE),
            (8.2, 1.65, "任务管理", "发布、指派、审核", WARN),
            (0.9, 4.05, "积分管理", "流水、统计、手动调整", GREEN),
            (4.55, 4.05, "成员管理", "角色、邀请码、访问状态", BRAND),
            (8.2, 4.05, "规则设置", "公告、点餐、审核规则", BLUE),
        ],
        1,
    ):
        card(s, x, y, 3.1, 1.45, h, b, accent=c, number=i)

    # 08 Menu
    s = add_slide(prs, 8, BRAND)
    title(s, "菜单系统", "菜单不是列表，是家庭餐桌的配置中心")
    bullet_list(s, ["分类管理：只看当前分类，避免无限展开。", "菜品编辑：名称、积分、排序、描述、图片。", "状态控制：上下架、今日推荐。", "静默筛选：切换分类不再重新进入页面。"], 0.95, 1.75, 6.9, 2.6, size=18, gap=0.58)
    card(s, 8.25, 1.75, 3.65, 3.65, "关键体验", "当分类越来越多，页面不能变成无尽长列表。\n现在只显示选中分类，管理更聚焦。", color=PANEL_2, accent=BRAND)

    # 09 Tasks
    s = add_slide(prs, 9, GREEN)
    title(s, "任务积分系统", "把家务变成有反馈的参与")
    steps = [("发布", "家主创建任务"), ("领取", "成员参与"), ("提交", "完成任务"), ("审核", "家主确认"), ("到账", "积分发放")]
    for i, (h, b) in enumerate(steps):
        card(s, 0.95 + i * 2.35, 2.35, 1.85, 1.45, h, b, accent=[BRAND, BLUE, WARN, GREEN, BRAND][i], number=i + 1)
        if i < 4:
            connector(s, 2.8 + i * 2.35, 3.07, 3.28 + i * 2.35, 3.07, BRAND)
    textbox(s, "完成一件事之后，系统马上承认这件事。反馈，是家庭协作持续下去的燃料。", Inches(1.25), Inches(5.35), Inches(10.8), Inches(0.6), size=20, color=MUTED, align=PP_ALIGN.CENTER)

    # 10 Permissions
    s = add_slide(prs, 10, BLUE)
    title(s, "权限与家庭邀请码", "一个家庭，一个空间，一个邀请码")
    card(s, 1.0, 1.8, 3.05, 2.9, "创建家庭", "家主创建家庭空间，成为默认 owner。", accent=BRAND, number=1)
    card(s, 4.95, 1.8, 3.05, 2.9, "邀请码加入", "成员注册后输入家庭邀请码完成绑定。", accent=BLUE, number=2)
    card(s, 8.9, 1.8, 3.05, 2.9, "角色管理", "家主维护成员角色和访问状态。", accent=WARN, number=3)
    textbox(s, "简单加入，清晰边界。后台管理端只给家主。", Inches(2.1), Inches(5.65), Inches(9.1), Inches(0.45), size=20, color=MUTED, align=PP_ALIGN.CENTER)

    # 11 UX
    s = add_slide(prs, 11, WARN)
    title(s, "体验升级", "从网页变成 App 感")
    for i, item in enumerate(["移动端底部导航", "自定义下拉弹层", "静默筛选", "暗色玻璃主题", "云端部署", "线上 e2e 回归"]):
        card(s, 0.95 + (i % 3) * 3.9, 1.6 + (i // 3) * 2.15, 3.25, 1.35, item, "更少跳转，更像一个真正能用的家庭工具。", accent=[BRAND, BLUE, WARN, GREEN, BRAND, BLUE][i])

    # 12 Architecture
    s = add_slide(prs, 12, BLUE)
    title(s, "技术架构", "轻量，但完整")
    card(s, 0.9, 2.45, 2.2, 1.15, "Browser", "Mobile / Desktop", accent=BRAND)
    card(s, 4.0, 1.55, 2.55, 1.15, "Web Worker", "家庭端", accent=BLUE)
    card(s, 4.0, 3.45, 2.55, 1.15, "Admin Worker", "家主管理端", accent=WARN)
    card(s, 8.0, 2.45, 3.0, 1.15, "Supabase", "Auth / Postgres / Storage", accent=GREEN)
    connector(s, 3.1, 3.02, 4.0, 2.12, BRAND)
    connector(s, 3.1, 3.02, 4.0, 4.02, BRAND)
    connector(s, 6.55, 2.12, 8.0, 3.02, BLUE)
    connector(s, 6.55, 4.02, 8.0, 3.02, WARN)
    textbox(s, "Next.js App Router + Cloudflare Workers + Supabase", Inches(2.5), Inches(5.65), Inches(8.5), Inches(0.45), size=20, color=MUTED, align=PP_ALIGN.CENTER)

    # 13 Cloud
    s = add_slide(prs, 13, GREEN)
    title(s, "云端闭环", "从代码到线上，每一步自动化")
    steps = [("GitHub", "提交代码"), ("Actions", "构建检查"), ("Cloudflare", "Web/Admin 部署"), ("Supabase", "数据与认证"), ("Playwright", "线上回归")]
    for i, (h, b) in enumerate(steps):
        card(s, 0.9 + i * 2.45, 2.45, 1.95, 1.35, h, b, accent=[BRAND, BLUE, WARN, GREEN, BRAND][i], number=i + 1)
        if i < 4:
            connector(s, 2.85 + i * 2.45, 3.12, 3.35 + i * 2.45, 3.12, LINE)
    textbox(s, "项目不是只在本地跑通，而是已经形成持续上线能力。", Inches(2.4), Inches(5.55), Inches(8.8), Inches(0.45), size=20, color=MUTED, align=PP_ALIGN.CENTER)

    # 14 Value
    s = add_slide(prs, 14, BRAND)
    title(s, "项目价值", "让家庭日常有一个温柔的系统")
    bullet_list(s, ["点餐更清楚。", "家务更有反馈。", "积分更有意义。", "家主更省心。", "家人更容易参与。"], 1.05, 1.75, 5.5, 3.1, size=22, gap=0.65)
    textbox(s, "好的家庭工具，不是让家庭变得机械，\n而是让协作更少摩擦。", Inches(7.0), Inches(2.2), Inches(4.9), Inches(1.3), size=28, bold=True)

    # 15 Roadmap
    s = add_slide(prs, 15, BLUE)
    title(s, "未来路线图", "更有生命力的家庭小宇宙")
    items = [("积分商城", "把积分变成更多家庭权益"), ("排行榜", "让参与可见但不制造压力"), ("成就徽章", "给任务完成更多仪式感"), ("消息提醒", "点餐和任务不再错过"), ("PWA", "像 App 一样安装和打开"), ("AI 推荐", "根据家庭偏好推荐菜单")]
    for i, (h, b) in enumerate(items):
        card(s, 0.95 + (i % 3) * 3.9, 1.6 + (i // 3) * 2.15, 3.25, 1.35, h, b, accent=[BRAND, BLUE, WARN, GREEN, BRAND, BLUE][i])

    # 16 End
    s = add_slide(prs, 16, WARN)
    textbox(s, "Night Food", Inches(0.85), Inches(1.35), Inches(4.3), Inches(0.45), size=23, color=BRAND, bold=True, font="Aptos")
    textbox(s, "家宴星球", Inches(0.78), Inches(2.0), Inches(5.5), Inches(0.95), size=56, bold=True)
    textbox(s, "让每一次点餐和每一次任务，\n都成为家庭协作的一部分。", Inches(0.85), Inches(3.25), Inches(7.2), Inches(0.95), size=25, color=MUTED)
    card(s, 8.25, 1.75, 3.6, 3.7, "发布完成", "核心闭环已上线\nWeb/Admin 云端部署\n线上 e2e 通过\n下一站：家庭小宇宙 2.0", color=PANEL_2, accent=BRAND)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"generated {OUT}")


if __name__ == "__main__":
    build()
